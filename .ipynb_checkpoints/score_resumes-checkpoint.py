
import os
import json
import anthropic
import google.generativeai as genai
from pathlib import Path
from typing import List, Dict, Any
import PyPDF2
import docx
import pandas as pd
from datetime import datetime
import argparse


class ResumeScorer:
    """
    Scores resumes using a pre-generated rubric with multiple LLMs.
    Outputs detailed scoring, rankings, and explanations.
    """
    
    def __init__(self, resume_dir: str = "resumes", rubric_path: str = "rubric.json"):
        """
        Initialize the resume scorer.
        
        Args:
            resume_dir: Directory containing resume files
            rubric_path: Path to the rubric JSON file
        """
        self.resume_dir = Path(resume_dir)
        self.rubric_path = Path(rubric_path)
        self.rubric = None
        self.resumes = []
        self.scores = []
        
        # Initialize API clients
        self.claude_client = anthropic.Anthropic(
            api_key=os.environ.get("ANTHROPIC_API_KEY")
        )
        
        genai.configure(api_key=os.environ.get("GOOGLE_API_KEY"))
        self.gemini_model = genai.GenerativeModel('gemini-2.0-flash-exp')
        
    def load_rubric(self):
        """Load the rubric from JSON file."""
        if not self.rubric_path.exists():
            raise FileNotFoundError(f"Rubric file '{self.rubric_path}' not found")
        
        with open(self.rubric_path, 'r') as f:
            self.rubric = json.load(f)
        
        print(f"✓ Loaded rubric from: {self.rubric_path}\n")
        return self.rubric
    
    def extract_text_from_pdf(self, file_path: Path) -> str:
        """Extract text from PDF file."""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text()
        except Exception as e:
            print(f"Error reading PDF {file_path}: {e}")
        return text
    
    def extract_text_from_docx(self, file_path: Path) -> str:
        """Extract text from DOCX file."""
        text = ""
        try:
            doc = docx.Document(file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        except Exception as e:
            print(f"Error reading DOCX {file_path}: {e}")
        return text
    
    def extract_text_from_txt(self, file_path: Path) -> str:
        """Extract text from TXT file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except Exception as e:
            print(f"Error reading TXT {file_path}: {e}")
            return ""
    
    def load_resumes(self) -> List[Dict[str, Any]]:
        """Load all resumes from the resume directory."""
        print(f"Loading resumes from {self.resume_dir}...")
        
        if not self.resume_dir.exists():
            raise FileNotFoundError(f"Resume directory '{self.resume_dir}' not found")
        
        resume_files = list(self.resume_dir.glob("*"))
        
        for file_path in resume_files:
            if file_path.is_file():
                text = ""
                
                if file_path.suffix.lower() == '.pdf':
                    text = self.extract_text_from_pdf(file_path)
                elif file_path.suffix.lower() == '.docx':
                    text = self.extract_text_from_docx(file_path)
                elif file_path.suffix.lower() == '.txt':
                    text = self.extract_text_from_txt(file_path)
                else:
                    print(f"Skipping unsupported file type: {file_path}")
                    continue
                
                if text.strip():
                    self.resumes.append({
                        'filename': file_path.name,
                        'content': text.strip(),
                        'file_path': str(file_path)
                    })
                    print(f"✓ Loaded: {file_path.name}")
        
        print(f"\nTotal resumes loaded: {len(self.resumes)}\n")
        return self.resumes
    
    def score_resume_with_llm(self, resume: Dict[str, Any], model: str = "claude") -> Dict[str, Any]:
        """
        Score a single resume using an LLM with granular decimal scoring.
        
        Args:
            resume: Resume dictionary with content
            model: Which LLM to use ("claude" or "gemini")
            
        Returns:
            Dictionary containing detailed scores and evaluation
        """
        prompt = f"""You are an expert evaluator for a VC+Founders Dinner event. Score this candidate's resume using the provided rubric.

**RUBRIC:**
{json.dumps(self.rubric, indent=2)}

**CANDIDATE RESUME:**
{resume['content']}

**CRITICAL SCORING INSTRUCTIONS:**

1. **Use DECIMAL PRECISION**: You MUST score using decimal points (e.g., 14.3, 7.8, 11.2). Do NOT use round numbers.

2. **Fine-Grained Differentiation**: Make subtle distinctions between candidates:
   - If a criterion has max 20 points, scores should vary across the full range (0.0 to 20.0)
   - Use at least one decimal place for every score
   - Avoid clustering scores around the same values
   - Be critical and discriminating - perfect scores (max points) should be extremely rare

3. **Scoring Guidelines Interpretation**:
   - "High" (80-100% of max): 
     * Top 10% = 95-100% of max points (e.g., 19.0-20.0 for a 20-point criterion)
     * Strong = 85-94% of max points (e.g., 17.0-18.8)
     * Good = 80-84% of max points (e.g., 16.0-16.9)
   - "Medium" (40-79% of max):
     * Upper-medium = 65-79% of max points (e.g., 13.0-15.8)
     * Mid-medium = 50-64% of max points (e.g., 10.0-12.8)
     * Lower-medium = 40-49% of max points (e.g., 8.0-9.9)
   - "Low" (0-39% of max):
     * Some evidence = 20-39% of max points (e.g., 4.0-7.8)
     * Minimal = 5-19% of max points (e.g., 1.0-3.8)
     * None = 0-4% of max points (e.g., 0.0-0.8)

4. **Evidence-Based Scoring**: 
   - Justify each decimal score with specific resume evidence
   - More evidence = higher within-band precision
   - Partial evidence = mid-to-lower within-band score

5. **Comparative Thinking**:
   - Consider how this candidate compares to theoretical ideal candidates
   - Reserve top scores for truly exceptional achievements
   - Use the full scoring spectrum

**OUTPUT FORMAT:**
Return your evaluation as a JSON object with this structure:
{{
  "crackedness_scores": [
    {{
      "criterion": "criterion name",
      "points_awarded": X.X,
      "max_points": Y,
      "percentage": Z.Z,
      "evidence": "specific evidence from resume justifying this exact score"
    }}
  ],
  "fit_scores": [
    {{
      "criterion": "criterion name",
      "points_awarded": X.X,
      "max_points": Y,
      "percentage": Z.Z,
      "evidence": "specific evidence from resume justifying this exact score"
    }}
  ],
  "total_crackedness": X.X,
  "total_fit": Y.Y,
  "candidate_description": "one paragraph description",
  "strengths_explanation": "why this candidate stands out or falls short"
}}

**REMEMBER**: Use decimal precision (e.g., 14.7, not 15.0). Differentiate carefully. Be critical.

Return ONLY valid JSON, no other text."""
        
        try:
            if model == "claude":
                response = self.claude_client.messages.create(
                    model="claude-sonnet-4-6",
                    max_tokens=8000,
                    temperature=0.4,  # Increased from 0.2 for more variability
                    messages=[{"role": "user", "content": prompt}]
                )
                response_text = response.content[0].text
            else:  # gemini
                response = self.gemini_model.generate_content(
                    prompt,
                    generation_config=genai.GenerationConfig(
                        temperature=0.4,  # Increased from 0.2
                        max_output_tokens=8000,
                    )
                )
                response_text = response.text
            
            # Extract JSON from response
            json_start = response_text.find('{')
            json_end = response_text.rfind('}') + 1
            if json_start != -1 and json_end > json_start:
                score_json = response_text[json_start:json_end]
                result = json.loads(score_json)
                
                # Add percentage calculations if not present
                for score_list in ['crackedness_scores', 'fit_scores']:
                    for item in result.get(score_list, []):
                        if 'percentage' not in item:
                            item['percentage'] = round((item['points_awarded'] / item['max_points']) * 100, 1)
                
                return result
            else:
                raise ValueError(f"No valid JSON found in {model} response")
                
        except Exception as e:
            print(f"Error scoring with {model}: {e}")
            return None
    
    def score_all_resumes(self, use_ensemble: bool = True):
        """
        Score all resumes using the rubric.
        
        Args:
            use_ensemble: If True, use both Claude and Gemini and average scores
        """
        if not self.rubric:
            raise ValueError("No rubric loaded. Call load_rubric() first.")
        
        if not self.resumes:
            raise ValueError("No resumes loaded. Call load_resumes() first.")
        
        print("\n" + "="*60)
        print("SCORING RESUMES")
        print("="*60 + "\n")
        
        for i, resume in enumerate(self.resumes, 1):
            print(f"Scoring {i}/{len(self.resumes)}: {resume['filename']}")
            
            if use_ensemble:
                # Score with both models
                claude_score = self.score_resume_with_llm(resume, model="claude")
                gemini_score = self.score_resume_with_llm(resume, model="gemini")
                
                if claude_score and gemini_score:
                    # Average the scores
                    final_score = {
                        'filename': resume['filename'],
                        'total_crackedness': round((claude_score['total_crackedness'] + gemini_score['total_crackedness']) / 2, 2),
                        'total_fit': round((claude_score['total_fit'] + gemini_score['total_fit']) / 2, 2),
                        'candidate_description': claude_score['candidate_description'],
                        'strengths_explanation': claude_score['strengths_explanation'],
                        'detailed_scores': {
                            'claude': claude_score,
                            'gemini': gemini_score
                        }
                    }
                elif claude_score:
                    final_score = {
                        'filename': resume['filename'],
                        'total_crackedness': round(claude_score['total_crackedness'], 2),
                        'total_fit': round(claude_score['total_fit'], 2),
                        'candidate_description': claude_score['candidate_description'],
                        'strengths_explanation': claude_score['strengths_explanation'],
                        'detailed_scores': {'claude': claude_score}
                    }
                elif gemini_score:
                    final_score = {
                        'filename': resume['filename'],
                        'total_crackedness': round(gemini_score['total_crackedness'], 2),
                        'total_fit': round(gemini_score['total_fit'], 2),
                        'candidate_description': gemini_score['candidate_description'],
                        'strengths_explanation': gemini_score['strengths_explanation'],
                        'detailed_scores': {'gemini': gemini_score}
                    }
                else:
                    print(f"  ✗ Failed to score {resume['filename']}")
                    continue
            else:
                # Use only Claude
                claude_score = self.score_resume_with_llm(resume, model="claude")
                if claude_score:
                    final_score = {
                        'filename': resume['filename'],
                        'total_crackedness': round(claude_score['total_crackedness'], 2),
                        'total_fit': round(claude_score['total_fit'], 2),
                        'candidate_description': claude_score['candidate_description'],
                        'strengths_explanation': claude_score['strengths_explanation'],
                        'detailed_scores': {'claude': claude_score}
                    }
                else:
                    print(f"  ✗ Failed to score {resume['filename']}")
                    continue
            
            self.scores.append(final_score)
            print(f"  ✓ Crackedness: {final_score['total_crackedness']:.2f}/100, Fit: {final_score['total_fit']:.2f}/100\n")
        
        print("="*60)
        print("SCORING COMPLETE")
        print("="*60 + "\n")
    
    def rank_candidates(self) -> List[Dict[str, Any]]:
        """
        Rank candidates based on their scores.
        Uses weighted combination: 60% Crackedness + 40% Fit
        
        Returns:
            List of candidates sorted by composite score
        """
        for score in self.scores:
            # Composite score: weighted combination
            score['composite_score'] = round((0.6 * score['total_crackedness']) + (0.4 * score['total_fit']), 2)
        
        # Sort by composite score (descending)
        ranked = sorted(self.scores, key=lambda x: x['composite_score'], reverse=True)
        
        # Add rank
        for i, score in enumerate(ranked, 1):
            score['rank'] = i
        
        return ranked
    
    def save_detailed_scores(self, output_dir: str = "rubric_scores"):
        """
        Save detailed scoring for each candidate to individual JSON files.
        
        Args:
            output_dir: Directory to save individual score files
        """
        output_path = Path(output_dir)
        output_path.mkdir(exist_ok=True)
        
        for score in self.scores:
            filename = score['filename'].rsplit('.', 1)[0] + '_score.json'
            filepath = output_path / filename
            
            with open(filepath, 'w') as f:
                json.dump(score, f, indent=2)
            
            print(f"✓ Saved detailed score: {filepath}")
        
        print(f"\n✓ All detailed scores saved to: {output_dir}/\n")
    
    def create_summary_spreadsheet(self, output_file: str = "candidate_rankings.xlsx"):
        """
        Create a summary spreadsheet with rankings and explanations.
        
        Args:
            output_file: Output Excel file path
        """
        ranked = self.rank_candidates()
        
        # Prepare data for spreadsheet
        data = []
        for candidate in ranked:
            data.append({
                'Rank': candidate['rank'],
                'Candidate': candidate['filename'],
                'Composite Score': candidate['composite_score'],
                'Crackedness Score': candidate['total_crackedness'],
                'Fit Score': candidate['total_fit'],
                'Description': candidate['candidate_description'],
                'Why Selected/Notable': candidate['strengths_explanation']
            })
        
        # Create DataFrame
        df = pd.DataFrame(data)
        
        # Save to Excel with formatting
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Rankings', index=False)
            
            # Get workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Rankings']
            
            # Adjust column widths
            worksheet.column_dimensions['A'].width = 8   # Rank
            worksheet.column_dimensions['B'].width = 30  # Candidate
            worksheet.column_dimensions['C'].width = 16  # Composite
            worksheet.column_dimensions['D'].width = 16  # Crackedness
            worksheet.column_dimensions['E'].width = 12  # Fit
            worksheet.column_dimensions['F'].width = 60  # Description
            worksheet.column_dimensions['G'].width = 60  # Why Selected
            
            # Enable text wrapping for description columns
            from openpyxl.styles import Alignment
            for row in worksheet.iter_rows(min_row=2, max_row=len(data)+1, min_col=6, max_col=7):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        print(f"✓ Summary spreadsheet saved to: {output_file}\n")
        
        return df
    
    def print_summary(self):
        """Print a summary of the scoring results."""
        ranked = self.rank_candidates()
        
        print("\n" + "="*60)
        print("CANDIDATE RANKINGS")
        print("="*60 + "\n")
        
        for candidate in ranked[:10]:  # Show top 10
            print(f"#{candidate['rank']} - {candidate['filename']}")
            print(f"   Composite: {candidate['composite_score']:.2f} | "
                  f"Crackedness: {candidate['total_crackedness']:.2f} | "
                  f"Fit: {candidate['total_fit']:.2f}")
            print(f"   {candidate['candidate_description'][:100]}...")
            print()
        
        if len(ranked) > 10:
            print(f"... and {len(ranked) - 10} more candidates\n")
        
        print("="*60 + "\n")


def main():
    """
    Main function to run the resume scorer with command-line arguments.
    """
    parser = argparse.ArgumentParser(
        description='Score resumes using a pre-generated rubric',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python score_resumes.py --rubric rubric.json --resume-dir resumes
  
  python score_resumes.py --rubric healthcare_rubric.json --output healthcare_rankings.xlsx
  
  python score_resumes.py --rubric rubric.json --no-ensemble --output-dir scored_resumes
        """
    )
    
    parser.add_argument(
        '--rubric',
        type=str,
        default='rubric.json',
        help='Path to rubric JSON file (default: rubric.json)'
    )
    
    parser.add_argument(
        '--resume-dir',
        type=str,
        default='resumes',
        help='Directory containing resumes (default: resumes)'
    )
    
    parser.add_argument(
        '--output',
        type=str,
        default='candidate_rankings.xlsx',
        help='Output Excel file for rankings (default: candidate_rankings.xlsx)'
    )
    
    parser.add_argument(
        '--output-dir',
        type=str,
        default='rubric_scores',
        help='Directory for detailed score JSON files (default: rubric_scores)'
    )
    
    parser.add_argument(
        '--no-ensemble',
        action='store_true',
        help='Use only Claude instead of ensemble (faster but less robust)'
    )
    
    args = parser.parse_args()
    
    print("\n" + "="*60)
    print("VC+FOUNDERS DINNER RESUME SCORER")
    print("="*60)
    print(f"\nRubric: {args.rubric}")
    print(f"Resume Directory: {args.resume_dir}")
    print(f"Output Spreadsheet: {args.output}")
    print(f"Detailed Scores Directory: {args.output_dir}")
    print(f"Ensemble Mode: {'No (Claude only)' if args.no_ensemble else 'Yes (Claude + Gemini)'}\n")
    
    # Initialize scorer
    scorer = ResumeScorer(resume_dir=args.resume_dir, rubric_path=args.rubric)
    
    # Load rubric and resumes
    scorer.load_rubric()
    scorer.load_resumes()
    
    # Score all resumes
    scorer.score_all_resumes(use_ensemble=not args.no_ensemble)
    
    # Save detailed scores
    scorer.save_detailed_scores(output_dir=args.output_dir)
    
    # Create summary spreadsheet
    scorer.create_summary_spreadsheet(output_file=args.output)
    
    # Print summary
    scorer.print_summary()
    
    print("Next steps:")
    print(f"1. Review candidate rankings in {args.output}")
    print(f"2. Check detailed scores in {args.output_dir}/ directory")
    print("3. Select candidates for the dinner based on rankings")


if __name__ == "__main__":
    main()
